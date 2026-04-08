import io
import csv as csv_lib
from typing import Union
from Providers.ai_provider import AIProvider

import fitz
from docx import Document
from pptx import Presentation


class FileExtractor:

    SUPPORTED_EXTENSIONS = {"pdf", "docx", "txt", "md", "csv", "pptx", "png", "jpg", "jpeg", "webp", "gif", "bmp"}
    MAX_FILE_SIZE = 10 * 1024 * 1024  # 10MB
    IMAGE_EXTENSIONS = {"png", "jpg", "jpeg", "webp", "gif", "bmp"}

    MIME_MAP = {
        "png": "image/png",
        "jpg": "image/jpeg",
        "jpeg": "image/jpeg",
        "webp": "image/webp",
        "gif": "image/gif",
        "bmp": "image/bmp",
    }

    def __init__(self, ai: AIProvider):
        self.ai = ai

    def is_image(self, filename: str) -> bool:
        """Check if a filename is an image type we handle as vision input."""
        ext = filename.lower().split(".")[-1]
        return ext in self.IMAGE_EXTENSIONS

    def get_image_mime(self, filename: str) -> str:
        """Return the MIME type for an image filename."""
        ext = filename.lower().split(".")[-1]
        return self.MIME_MAP.get(ext, "image/png")

    async def extract_text(self, file_bytes: bytes, filename: str) -> str:
        """
        Extract text from a document file. Raises ValueError for images —
        images should be handled directly by the chat call instead of being
        converted to text via a separate Gemini call.
        
        Call is_image(filename) first to decide which path to take.
        """
        if len(file_bytes) > self.MAX_FILE_SIZE:
            raise ValueError(f"File exceeds 10MB limit ({len(file_bytes) / 1024 / 1024:.1f}MB)")

        ext = filename.lower().split(".")[-1]

        if ext not in self.SUPPORTED_EXTENSIONS:
            raise ValueError(f"Unsupported file type: .{ext}")

        if ext in self.IMAGE_EXTENSIONS:
            raise ValueError(
                "Image files should be passed directly to the chat call, "
                "not extracted to text. Use is_image() to detect images first."
            )

        if ext == "pdf":
            return self._extract_pdf(file_bytes)
        elif ext == "docx":
            return self._extract_docx(file_bytes)
        elif ext in ("txt", "md"):
            return file_bytes.decode("utf-8", errors="ignore")
        elif ext == "csv":
            return self._extract_csv(file_bytes)
        elif ext == "pptx":
            return self._extract_pptx(file_bytes)

    def _extract_pdf(self, file_bytes: bytes) -> str:
        doc = fitz.open(stream=file_bytes, filetype="pdf")
        pages = []
        for i, page in enumerate(doc):
            text = page.get_text().strip()
            if text:
                pages.append(f"[Page {i + 1}]\n{text}")
        doc.close()
        if not pages:
            raise ValueError("PDF has no extractable text. Try uploading as an image instead.")
        return "\n\n".join(pages)

    def _extract_docx(self, file_bytes: bytes) -> str:
        doc = Document(io.BytesIO(file_bytes))
        chunks = []
        for para in doc.paragraphs:
            if para.text.strip():
                chunks.append(para.text.strip())
        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                if row_text:
                    chunks.append(row_text)
        return "\n\n".join(chunks)

    def _extract_csv(self, file_bytes: bytes) -> str:
        decoded = file_bytes.decode("utf-8", errors="ignore")
        reader = csv_lib.DictReader(io.StringIO(decoded))
        rows = []
        for row in reader:
            row_text = ", ".join(f"{k}: {v}" for k, v in row.items() if v)
            if row_text:
                rows.append(row_text)
        return "\n".join(rows)

    def _extract_pptx(self, file_bytes: bytes) -> str:
        prs = Presentation(io.BytesIO(file_bytes))
        slides = []
        for i, slide in enumerate(prs.slides):
            texts = [shape.text.strip() for shape in slide.shapes if hasattr(shape, "text") and shape.text.strip()]
            if texts:
                slides.append(f"[Slide {i + 1}]\n" + "\n".join(texts))
        return "\n\n".join(slides)